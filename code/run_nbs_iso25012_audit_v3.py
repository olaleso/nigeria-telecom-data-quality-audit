
from pathlib import Path
import csv
import json
import math
import re
from collections import Counter, defaultdict

import openpyxl

ROOT = Path(r"C:\PhDDataSets\NBS")
EXTRACTED = ROOT / "extracted"
OUT = ROOT / "iso25012_audit_v3"
OUT.mkdir(parents=True, exist_ok=True)

COUNT_TOL = 0.5
PCT_TOL = 0.01
ZONE_TOL = 0.5

EXPECTED_STATES = [
    "ABIA","ADAMAWA","AKWA IBOM","ANAMBRA","BAUCHI","BAYELSA","BENUE","BORNO",
    "CROSS RIVER","DELTA","EBONYI","EDO","EKITI","ENUGU","FCT","GOMBE","IMO",
    "JIGAWA","KADUNA","KANO","KATSINA","KEBBI","KOGI","KWARA","LAGOS","NASARAWA",
    "NIGER","OGUN","ONDO","OSUN","OYO","PLATEAU","RIVERS","SOKOTO","TARABA","YOBE",
    "ZAMFARA"
]

STATE_ALIASES = {
    "NASSARAWA": "NASARAWA",
    "FCT ABUJA": "FCT",
    "ABUJA FCT": "FCT",
    "FEDERAL CAPITAL TERRITORY": "FCT",
    "FEDERAL CAPITAL TERRITORY ABUJA": "FCT",
}

ZONE_MAP = {
    "ABIA":"SOUTH EAST","ADAMAWA":"NORTH EAST","AKWA IBOM":"SOUTH SOUTH",
    "ANAMBRA":"SOUTH EAST","BAUCHI":"NORTH EAST","BAYELSA":"SOUTH SOUTH",
    "BENUE":"NORTH CENTRAL","BORNO":"NORTH EAST","CROSS RIVER":"SOUTH SOUTH",
    "DELTA":"SOUTH SOUTH","EBONYI":"SOUTH EAST","EDO":"SOUTH SOUTH","EKITI":"SOUTH WEST",
    "ENUGU":"SOUTH EAST","FCT":"NORTH CENTRAL","GOMBE":"NORTH EAST","IMO":"SOUTH EAST",
    "JIGAWA":"NORTH WEST","KADUNA":"NORTH WEST","KANO":"NORTH WEST","KATSINA":"NORTH WEST",
    "KEBBI":"NORTH WEST","KOGI":"NORTH CENTRAL","KWARA":"NORTH CENTRAL","LAGOS":"SOUTH WEST",
    "NASARAWA":"NORTH CENTRAL","NIGER":"NORTH CENTRAL","OGUN":"SOUTH WEST","ONDO":"SOUTH WEST",
    "OSUN":"SOUTH WEST","OYO":"SOUTH WEST","PLATEAU":"NORTH CENTRAL","RIVERS":"SOUTH SOUTH",
    "SOKOTO":"NORTH WEST","TARABA":"NORTH EAST","YOBE":"NORTH EAST","ZAMFARA":"NORTH WEST"
}
EXPECTED_ZONES = ["NORTH CENTRAL","NORTH EAST","NORTH WEST","SOUTH EAST","SOUTH SOUTH","SOUTH WEST"]

KNOWN_STATE_LABEL_ISSUES = {"NASSARAWA": "NASARAWA"}

issues = []
sheet_occ = []
state_records = []
state_metrics = []
porting_records = []
zone_records = []
formula_records = []
metadata_records = []
schema_records = []

def txt(v):
    return "" if v is None else re.sub(r"\s+", " ", str(v).strip())

def up(v):
    s = txt(v).upper().replace("&", "AND")
    s = re.sub(r"[^A-Z0-9]+", " ", s)
    return re.sub(r"\s+", " ", s).strip()

def num(v):
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        if isinstance(v, float) and math.isnan(v):
            return None
        return float(v)
    s = str(v).strip().replace(",", "")
    if s in {"", "-", "–", "—", "NA", "N/A"}:
        return None
    try:
        return float(s)
    except Exception:
        return None

def norm_state(v):
    s = up(v)
    return STATE_ALIASES.get(s, s)

def parse_q(text):
    # accepts Q4 2025, Q4_2025, Q4, 2025, Q4-2025, etc.
    m = re.search(r"(?<![A-Za-z0-9])Q\s*([1-4])[\s,_\-:/]*?(20\d{2})\b", str(text), re.I)
    if not m:
        return None
    q, y = int(m.group(1)), int(m.group(2))
    return f"Q{q} {y}", y * 10 + q

def parse_release(filename):
    return parse_q(Path(filename).stem)

def stype(sheet):
    u = sheet.upper()
    if "VOICE" in u: return "Voice"
    if "INTERNET" in u: return "Internet"
    if "PORTING" in u: return "Porting"
    return "Other"

def is_serial(v):
    x = num(v)
    return x is not None and 1 <= x <= 100 and abs(x-round(x)) < 1e-9

def add_issue(severity, dimension, workbook, sheet, quarter, issue_type, detail,
              observed="", expected="", state="", cell="", classification="candidate"):
    issues.append({
        "severity": severity,
        "dimension": dimension,
        "workbook": workbook,
        "sheet": sheet,
        "quarter": quarter,
        "issue_type": issue_type,
        "state_or_group": state,
        "cell": cell,
        "observed": observed,
        "expected": expected,
        "classification": classification,
        "detail": detail,
    })

def find_row(ws, predicate, max_rows=25):
    for r in range(1, min(ws.max_row, max_rows)+1):
        vals = [ws.cell(r,c).value for c in range(1, ws.max_column+1)]
        if predicate(vals):
            return r
    return None

def find_operator_row(ws, service_type):
    if service_type not in {"Voice","Internet"}:
        return None
    def pred(vals):
        u = [up(x) for x in vals]
        hits = sum(k in u for k in ["MTN","GLO","AIRTEL","EMTS"])
        return hits >= 3 and any(x in {"SUB TOTAL","SUBTOTAL"} for x in u)
    return find_row(ws, pred, 20)

def find_subtotal_col(ws, operator_row):
    if not operator_row: return None
    for c in range(1, ws.max_column+1):
        if up(ws.cell(operator_row,c).value) in {"SUB TOTAL","SUBTOTAL"}:
            return c
    return None

def find_state_rows(ws):
    return [r for r in range(1, ws.max_row+1)
            if is_serial(ws.cell(r,1).value) and txt(ws.cell(r,2).value)]

def find_reference_columns(ws, service_type, operator_row):
    refs=[]
    if not operator_row: return refs
    for r in range(1, operator_row):
        for c in range(1, ws.max_column+1):
            label=txt(ws.cell(r,c).value)
            if service_type.upper() not in label.upper():
                continue
            pq=parse_q(label)
            if not pq:
                continue
            pct_col=None
            # Usually next column is YoY/QoQ
            for cc in range(c+1, min(c+3, ws.max_column+1)):
                if up(ws.cell(r,cc).value) in {"YOY","QOQ"}:
                    pct_col=cc
                    break
            refs.append({
                "row":r, "reference_col":c, "reference_quarter":pq[0],
                "pct_col":pct_col, "pct_label":up(ws.cell(r,pct_col).value) if pct_col else ""
            })
    # dedupe by reference_col
    seen=set(); out=[]
    for x in refs:
        if x["reference_col"] not in seen:
            out.append(x); seen.add(x["reference_col"])
    return out

def sheet_title_blob(ws, rows=12):
    vals=[]
    for r in range(1, min(ws.max_row,rows)+1):
        for c in range(1, ws.max_column+1):
            v=txt(ws.cell(r,c).value)
            if v: vals.append(v)
    return " | ".join(vals)

def period_match(ws, quarter):
    blob = sheet_title_blob(ws)
    if quarter in blob.upper().replace(",",""):
        return True
    q=int(quarter[1]); year=quarter.split()[1]
    month={1:"MARCH",2:"JUNE",3:"SEPTEMBER",4:"DECEMBER"}[q]
    return year in blob and month in blob.upper()

def source_note_scan(ws):
    found=[]
    for row in ws.iter_rows():
        for cell in row:
            u=up(cell.value)
            if any(k in u for k in ["SOURCE","NIGERIAN COMMUNICATIONS COMMISSION","NCC","NOTE"]):
                found.append(f"{cell.coordinate}: {txt(cell.value)}")
    return found[:20]

def count_formulas(ws):
    return sum(1 for row in ws.iter_rows() for cell in row
               if isinstance(cell.value,str) and cell.value.startswith("="))

def find_porting_header(ws):
    for r in range(1, min(ws.max_row,25)+1):
        vals=[up(ws.cell(r,c).value) for c in range(1,ws.max_column+1)]
        if "OPERATOR" in vals and "TOTAL" in vals:
            op_col=vals.index("OPERATOR")+1
            total_col=vals.index("TOTAL")+1
            return r,op_col,total_col
    return None,None,None

def find_zone_header(ws):
    for r in range(1, min(ws.max_row,30)+1):
        vals=[up(ws.cell(r,c).value) for c in range(1,ws.max_column+1)]
        if "INTERNET SUBSCRIPTION" in vals and "VOICE SUBSCRIPTION" in vals:
            ic=vals.index("INTERNET SUBSCRIPTION")+1
            vc=vals.index("VOICE SUBSCRIPTION")+1
            return r,ic,vc
    return None,None,None

def nearest_zone_col_left(ws, row, service_col):
    for c in range(service_col-1,0,-1):
        v=up(ws.cell(row,c).value)
        if v:
            return c
    return None

def write_csv(path, rows):
    if not rows:
        path.write_text("",encoding="utf-8")
        return
    keys=[]
    for r in rows:
        for k in r:
            if k not in keys: keys.append(k)
    with path.open("w",newline="",encoding="utf-8-sig") as f:
        w=csv.DictWriter(f,fieldnames=keys)
        w.writeheader(); w.writerows(rows)

xlsx=sorted(EXTRACTED.rglob("*.xlsx"))
if not xlsx:
    raise SystemExit(f"No .xlsx workbooks found under {EXTRACTED}")

books=[]
skipped_workbooks=[]
for p in xlsx:
    rel=parse_release(p.name)
    if not rel:
        skipped_workbooks.append(p.name)
        print("Skipping unparseable filename:",p.name)
        continue
    print("Loading:",p.name)
    books.append({
        "name":p.name,"release_q":rel[0],"release_key":rel[1],
        "wf":openpyxl.load_workbook(p,data_only=False),
        "wv":openpyxl.load_workbook(p,data_only=True),
    })
books.sort(key=lambda x:x["release_key"])

print(f"Discovered .xlsx files: {len(xlsx)}")
print(f"Parsed workbooks:       {len(books)}")
if skipped_workbooks:
    print("WARNING - skipped workbook names:")
    for name in skipped_workbooks:
        print("  -", name)

# This study corpus is expected to contain seven NBS quarterly workbooks.
# Do not silently continue with a partial corpus.
if len(books) != 7:
    raise SystemExit(
        f"\nSTOP: expected 7 parsed NBS workbooks but parsed {len(books)}. "
        f"Discovered {len(xlsx)} .xlsx files. "
        "Check the skipped names printed above before using any audit results."
    )

for b in books:
    for s in b["wf"].sheetnames:
        q=parse_q(s); typ=stype(s)
        if not q or typ=="Other": continue
        quarter,qkey=q
        wf=b["wf"][s]; wv=b["wv"][s]

        sheet_occ.append({
            "workbook":b["name"],"release_quarter":b["release_q"],
            "release_key":b["release_key"],"sheet":s,"quarter":quarter,
            "quarter_key":qkey,"service_type":typ
        })

        fcount=count_formulas(wf)
        formula_records.append({
            "workbook":b["name"],"release_quarter":b["release_q"],"sheet":s,
            "quarter":quarter,"service_type":typ,"formula_count":fcount
        })

        notes=source_note_scan(wf)
        pm=period_match(wf,quarter)
        metadata_records.append({
            "workbook":b["name"],"release_quarter":b["release_q"],"sheet":s,
            "quarter":quarter,"service_type":typ,
            "period_title_match":pm,
            "embedded_source_or_note_found":bool(notes),
            "examples":" || ".join(notes)
        })

        if not pm:
            add_issue("Medium","Understandability",b["name"],s,quarter,
                      "PERIOD_TITLE_NOT_DETECTED",
                      "Reporting period was not detected reliably in the sheet title area.",
                      classification="needs_manual_review")

        if typ in {"Voice","Internet"}:
            op_row=find_operator_row(wf,typ)
            sub_col=find_subtotal_col(wf,op_row)
            rows=find_state_rows(wf)
            refs=find_reference_columns(wf,typ,op_row)

            if op_row is None or sub_col is None:
                add_issue("High","Traceability",b["name"],s,quarter,
                          "STATE_TABLE_STRUCTURE_NOT_PARSED",
                          "Operator header/subtotal structure could not be identified dynamically.",
                          classification="parser_warning")
                continue

            operators=[]
            for c in range(3,sub_col):
                label=txt(wf.cell(op_row,c).value)
                operators.append((c,label or f"COL_{c}"))

            raw_states=[up(wf.cell(r,2).value) for r in rows]
            norm_states=[norm_state(wf.cell(r,2).value) for r in rows]

            missing=sorted(set(EXPECTED_STATES)-set(norm_states))
            extra=sorted(set(norm_states)-set(EXPECTED_STATES))
            duplicates=sorted(x for x,n in Counter(norm_states).items() if n>1)

            if len(rows)!=37 or missing:
                add_issue("High","Completeness",b["name"],s,quarter,
                          "STATE_COVERAGE",
                          "State/FCT coverage is incomplete after standardised-label normalisation.",
                          observed=f"rows={len(rows)}; missing={' | '.join(missing)}",
                          expected="37 states/FCT",classification="confirmed_if_persists")

            if extra or duplicates:
                add_issue("Medium","Consistency",b["name"],s,quarter,
                          "STATE_KEY_INCONSISTENCY",
                          "Unexpected or duplicate normalised state keys were found.",
                          observed=f"extra={extra}; duplicates={duplicates}",
                          classification="confirmed_if_persists")

            for raw,norm in zip(raw_states,norm_states):
                if raw != norm and raw in KNOWN_STATE_LABEL_ISSUES:
                    add_issue("Medium","Understandability",b["name"],s,quarter,
                              "NONSTANDARD_STATE_LABEL",
                              "A non-standard state spelling is used but maps unambiguously to the expected state.",
                              observed=raw, expected=norm, state=norm,
                              classification="confirmed_governance_issue")

            arithmetic_tests=0; arithmetic_pass=0
            pct_tests=0; pct_pass=0
            nonint=0; blank_ct=0; hyphen_ct=0; zero_ct=0

            max_expected_col=max([sub_col]+[
                x["pct_col"] or x["reference_col"] for x in refs
            ])

            for r in rows:
                raw_state=up(wf.cell(r,2).value)
                state=norm_state(wf.cell(r,2).value)
                operator_vals={}
                sum_ops=0.0
                for c,label in operators:
                    raw=wf.cell(r,c).value
                    v=num(wv.cell(r,c).value)
                    if txt(raw) in {"-","–","—"}: hyphen_ct+=1
                    elif raw in (None,""): blank_ct+=1
                    elif v==0: zero_ct+=1
                    if v is not None:
                        sum_ops+=v
                        if abs(v-round(v))>1e-9: nonint+=1
                    operator_vals[label]=v

                total=num(wv.cell(r,sub_col).value)
                if total is None:
                    # formula cache missing -> recompute, but record provenance
                    total=sum_ops
                    total_source="recomputed_from_operator_components"
                else:
                    total_source="published_or_cached"

                arithmetic_tests+=1
                diff=total-sum_ops
                ok=abs(diff)<=COUNT_TOL
                arithmetic_pass+=int(ok)
                if not ok:
                    add_issue("High","Accuracy / reconciliation",b["name"],s,quarter,
                              "STATE_SUBTOTAL_MISMATCH",
                              "State subtotal does not reconcile to operator components.",
                              observed=total,expected=sum_ops,state=state,
                              classification="confirmed_numeric_issue")

                rec={
                    "workbook":b["name"],"release_quarter":b["release_q"],
                    "release_key":b["release_key"],"sheet":s,"quarter":quarter,
                    "quarter_key":qkey,"service_type":typ,"row":r,
                    "raw_state_label":raw_state,"state":state,"current_total":total,
                    "current_total_source":total_source,"operator_sum":sum_ops,
                    "subtotal_difference":diff,"subtotal_pass":ok
                }
                for k,v in operator_vals.items():
                    rec[f"operator::{k}"]=v

                for ref in refs:
                    rv=num(wv.cell(r,ref["reference_col"]).value)
                    pv=num(wv.cell(r,ref["pct_col"]).value) if ref["pct_col"] else None
                    rec[f"reference::{ref['reference_quarter']}"]=rv
                    if ref["pct_col"]:
                        rec[f"pct::{ref['pct_label']}::{ref['reference_quarter']}"]=pv
                    if rv not in (None,0) and pv is not None:
                        expected_pct=(total-rv)/rv*100.0
                        pct_tests+=1
                        okp=abs(pv-expected_pct)<=PCT_TOL
                        pct_pass+=int(okp)
                        if not okp:
                            add_issue("High","Accuracy / reconciliation",b["name"],s,quarter,
                                      "PERCENTAGE_MISMATCH",
                                      f"{ref['pct_label']} does not reconcile to the current and reference values.",
                                      observed=pv,expected=expected_pct,state=state,
                                      classification="confirmed_numeric_issue")

                # Flag populated cells after expected reporting structure.
                for c in range(max_expected_col+1,wf.max_column+1):
                    raw=wf.cell(r,c).value
                    if raw not in (None,""):
                        add_issue("Low","Understandability",b["name"],s,quarter,
                                  "UNLABELLED_EXTRA_VALUE",
                                  "A populated cell occurs beyond the expected reporting columns.",
                                  observed=raw,state=state,cell=wf.cell(r,c).coordinate,
                                  classification="confirmed_presentation_issue")

                state_records.append(rec)

            state_metrics.append({
                "workbook":b["name"],"release_quarter":b["release_q"],"sheet":s,
                "quarter":quarter,"service_type":typ,
                "state_rows":len(rows),
                "normalised_state_coverage_pct":100*len(set(norm_states)&set(EXPECTED_STATES))/37,
                "missing_state_count":len(missing),
                "duplicate_state_count":len(duplicates),
                "operator_column_count":len(operators),
                "state_subtotal_pass_pct":100*arithmetic_pass/arithmetic_tests if arithmetic_tests else "",
                "percentage_recalc_pass_pct":100*pct_pass/pct_tests if pct_tests else "",
                "noninteger_count_cells":nonint,
                "blank_operator_cells":blank_ct,"hyphen_operator_cells":hyphen_ct,
                "zero_operator_cells":zero_ct,
                "formula_count":fcount
            })

            schema_records.append({
                "workbook":b["name"],"release_quarter":b["release_q"],
                "quarter":quarter,"service_type":typ,
                "schema":" | ".join(label for _,label in operators)
            })

        elif typ=="Porting":
            hr,opcol,totalcol=find_porting_header(wf)
            zr, internet_col, voice_col=find_zone_header(wf)
            if hr is None:
                add_issue("High","Traceability",b["name"],s,quarter,
                          "PORTING_HEADER_NOT_PARSED",
                          "Porting operator header row was not identified.",
                          classification="parser_warning")
                continue

            operators=[]
            for c in range(opcol+1,totalcol):
                lab=txt(wf.cell(hr,c).value)
                if lab: operators.append((c,lab))

            schema_records.append({
                "workbook":b["name"],"release_quarter":b["release_q"],
                "quarter":quarter,"service_type":typ,
                "schema":" | ".join(l for _,l in operators)
            })

            direction_rows={}
            for r in range(hr+1,min(wf.max_row,hr+12)+1):
                d=up(wf.cell(r,opcol).value)
                if d in {"PORT IN","PORT OUT"}:
                    direction_rows[d]=r

            for d,r in direction_rows.items():
                comp=sum((num(wv.cell(r,c).value) or 0.0) for c,_ in operators)
                pub=num(wv.cell(r,totalcol).value)
                if pub is None: pub=comp
                diff=pub-comp
                ok=abs(diff)<=COUNT_TOL
                porting_records.append({
                    "workbook":b["name"],"release_quarter":b["release_q"],"sheet":s,
                    "quarter":quarter,"direction":d,"published_total":pub,
                    "recomputed_total":comp,"difference":diff,"pass":ok,
                    "operator_schema":" | ".join(l for _,l in operators)
                })
                if not ok:
                    add_issue("High","Accuracy / reconciliation",b["name"],s,quarter,
                              "PORTING_TOTAL_MISMATCH",
                              f"{d} total does not reconcile to operator components.",
                              observed=pub,expected=comp,classification="confirmed_numeric_issue")

            if "PORT IN" in direction_rows and "PORT OUT" in direction_rows:
                rin=direction_rows["PORT IN"]; rout=direction_rows["PORT OUT"]
                pin=num(wv.cell(rin,totalcol).value)
                pout=num(wv.cell(rout,totalcol).value)
                if pin is not None and pout is not None and abs(pin-pout)>COUNT_TOL:
                    add_issue("High","Consistency",b["name"],s,quarter,
                              "PORT_IN_OUT_IMBALANCE",
                              "National port-in and port-out totals differ.",
                              observed=f"port-in={pin}; port-out={pout}",
                              expected="equal national totals",
                              classification="confirmed_if_semantically_required")

            if zr is not None:
                izone_col=nearest_zone_col_left(wf,zr,internet_col)
                vzone_col=nearest_zone_col_left(wf,zr,voice_col)
                left_header=txt(wf.cell(zr,izone_col).value) if izone_col else ""
                right_header=txt(wf.cell(zr,vzone_col).value) if vzone_col else ""
                if up(left_header)!="ZONE" or up(right_header)!="ZONE":
                    add_issue("Low","Understandability",b["name"],s,quarter,
                              "ZONE_HEADER_SPELLING",
                              "Zone header is inconsistent or misspelled.",
                              observed=f"left={left_header}; right={right_header}",
                              expected="Zone | Zone",
                              classification="confirmed_presentation_issue")
                for r in range(zr+1,min(wf.max_row,zr+10)+1):
                    iz=up(wf.cell(r,izone_col).value) if izone_col else ""
                    vz=up(wf.cell(r,vzone_col).value) if vzone_col else ""
                    if iz in EXPECTED_ZONES:
                        zone_records.append({
                            "workbook":b["name"],"release_quarter":b["release_q"],
                            "release_key":b["release_key"],"sheet":s,"quarter":quarter,
                            "zone":iz,"service_type":"Internet",
                            "published_zone_total":num(wv.cell(r,internet_col).value)
                        })
                    if vz in EXPECTED_ZONES:
                        zone_records.append({
                            "workbook":b["name"],"release_quarter":b["release_q"],
                            "release_key":b["release_key"],"sheet":s,"quarter":quarter,
                            "zone":vz,"service_type":"Voice",
                            "published_zone_total":num(wv.cell(r,voice_col).value)
                        })

# Canonical = latest available copy for each quarter/service.
latest={}
for x in sheet_occ:
    k=(x["quarter"],x["service_type"])
    if k not in latest or x["release_key"]>latest[k]["release_key"]:
        latest[k]=x

canonical_state=[
    r for r in state_records
    if latest[(r["quarter"],r["service_type"])]["workbook"]==r["workbook"]
]

# Cross-period reference reconciliation
actual={(r["quarter"],r["service_type"],r["state"]):r["current_total"] for r in canonical_state}
cross=[]
for r in canonical_state:
    for k,v in r.items():
        if not k.startswith("reference::") or v in (None,""): continue
        rq=k.split("::",1)[1]
        av=actual.get((rq,r["service_type"],r["state"]))
        if av is None: continue
        diff=float(v)-float(av)
        ok=abs(diff)<=COUNT_TOL
        cross.append({
            "quarter":r["quarter"],"service_type":r["service_type"],"state":r["state"],
            "reference_quarter":rq,"published_reference_value":v,
            "actual_reference_quarter_value":av,"difference":diff,"pass":ok
        })
        if not ok:
            add_issue("High","Consistency",r["workbook"],r["sheet"],r["quarter"],
                      "CROSS_PERIOD_REFERENCE_MISMATCH",
                      "Embedded prior-period value differs from the value in the referenced quarter dataset.",
                      observed=v,expected=av,state=r["state"],
                      classification="confirmed_numeric_issue")

# Zone reconciliation using normalised state labels.
state_zone=defaultdict(float)
for r in canonical_state:
    z=ZONE_MAP.get(r["state"])
    if z:
        state_zone[(r["quarter"],r["service_type"],z)]+=float(r["current_total"])

canonical_zone=[
    z for z in zone_records
    if latest.get((z["quarter"],"Porting"),{}).get("workbook")==z["workbook"]
]
zone_recon=[]
for z in canonical_zone:
    exp=state_zone.get((z["quarter"],z["service_type"],z["zone"]))
    pub=z["published_zone_total"]
    diff=None if exp is None or pub is None else pub-exp
    ok=None if diff is None else abs(diff)<=ZONE_TOL
    row=dict(z); row.update({"recomputed_from_states":exp,"difference":diff,"pass":ok})
    zone_recon.append(row)
    if ok is False:
        add_issue("High","Accuracy / reconciliation",z["workbook"],z["sheet"],z["quarter"],
                  "ZONE_TOTAL_MISMATCH",
                  f"{z['service_type']} zone total does not reconcile to the normalised state-level total.",
                  observed=pub,expected=exp,state=z["zone"],
                  classification="confirmed_numeric_issue")

# Revision stability
revision=[]
grouped=defaultdict(list)
for r in state_records:
    grouped[(r["quarter"],r["service_type"],r["state"])].append(r)
for k,grp in grouped.items():
    if len(grp)<2: continue
    grp=sorted(grp,key=lambda x:x["release_key"])
    first,last=grp[0],grp[-1]
    diff=float(last["current_total"])-float(first["current_total"])
    revision.append({
        "quarter":k[0],"service_type":k[1],"state":k[2],
        "earliest_workbook":first["workbook"],"latest_workbook":last["workbook"],
        "earliest_total":first["current_total"],"latest_total":last["current_total"],
        "difference":diff,"status":"same" if abs(diff)<=COUNT_TOL else "changed"
    })

# Formula transparency drift
formula_drift=[]
fg=defaultdict(list)
for r in formula_records:
    key=(r["quarter"],r["service_type"])
    rr=dict(r)
    rr["release_key"]=parse_release(r["workbook"])[1]
    fg[key].append(rr)
for key,grp in fg.items():
    if len(grp)<2: continue
    grp=sorted(grp,key=lambda x:x["release_key"])
    formula_drift.append({
        "quarter":key[0],"service_type":key[1],
        "earliest_workbook":grp[0]["workbook"],"latest_workbook":grp[-1]["workbook"],
        "earliest_formula_count":grp[0]["formula_count"],
        "latest_formula_count":grp[-1]["formula_count"],
        "formula_count_change":grp[-1]["formula_count"]-grp[0]["formula_count"]
    })

# Schema evolution based on canonical latest copies.
canonical_schema=[]
for r in schema_records:
    if latest.get((r["quarter"],r["service_type"]),{}).get("workbook")==r["workbook"]:
        canonical_schema.append(r)

schema_evo=[]
for typ in ["Voice","Internet","Porting"]:
    grp=[r for r in canonical_schema if r["service_type"]==typ]
    grp.sort(key=lambda x:parse_q(x["quarter"])[1])
    prev=None
    for r in grp:
        schema_evo.append({
            "quarter":r["quarter"],"service_type":typ,"workbook":r["workbook"],
            "schema":r["schema"],
            "changed_from_previous_quarter":False if prev is None else r["schema"]!=prev,
            "previous_schema":"" if prev is None else prev
        })
        prev=r["schema"]

# Summary
canon_metrics=[]
for m in state_metrics:
    if latest.get((m["quarter"],m["service_type"]),{}).get("workbook")==m["workbook"]:
        canon_metrics.append(m)

def pct_pass(rows, field="pass"):
    good=[r for r in rows if r.get(field) not in (None,"")]
    if not good: return None
    return 100*sum(str(r[field]).lower()=="true" or r[field] is True for r in good)/len(good)

state_arith=[float(r["state_subtotal_pass_pct"]) for r in canon_metrics if r["state_subtotal_pass_pct"]!=""]
pct_arith=[float(r["percentage_recalc_pass_pct"]) for r in canon_metrics if r["percentage_recalc_pass_pct"]!=""]
coverage=[float(r["normalised_state_coverage_pct"]) for r in canon_metrics]

port_canon=[r for r in porting_records if latest.get((r["quarter"],"Porting"),{}).get("workbook")==r["workbook"]]
zone_ok=[r for r in zone_recon if r["pass"] is not None]

summary=[
    {
        "dimension":"Completeness",
        "measure":"Mean normalised state/FCT coverage across canonical Voice/Internet sheets",
        "result":f"{sum(coverage)/len(coverage):.2f}%" if coverage else "Not available",
        "interpretation":"Label aliases are normalised before coverage is calculated; naming defects are reported separately."
    },
    {
        "dimension":"Consistency",
        "measure":"Cross-period embedded-reference reconciliation",
        "result":f"{pct_pass(cross):.2f}% ({sum(str(r['pass']).lower()=='true' or r['pass'] is True for r in cross)}/{len(cross)})" if cross else "Not available",
        "interpretation":"Tests published-table agreement, not independent real-world truth."
    },
    {
        "dimension":"Accuracy / reconciliation",
        "measure":"State subtotal arithmetic",
        "result":f"{sum(state_arith)/len(state_arith):.2f}%" if state_arith else "Not available",
        "interpretation":"Current state totals vs sum of operator components."
    },
    {
        "dimension":"Accuracy / reconciliation",
        "measure":"Published YoY/QoQ percentage recalculation",
        "result":f"{sum(pct_arith)/len(pct_arith):.2f}%" if pct_arith else "Not available",
        "interpretation":"Recomputes percentages where both current and reference values are available."
    },
    {
        "dimension":"Accuracy / reconciliation",
        "measure":"Porting component-to-total reconciliation",
        "result":f"{pct_pass(port_canon):.2f}% ({sum(str(r['pass']).lower()=='true' or r['pass'] is True for r in port_canon)}/{len(port_canon)})" if port_canon else "Not available",
        "interpretation":"PORT-IN/PORT-OUT totals vs operator components."
    },
    {
        "dimension":"Accuracy / reconciliation",
        "measure":"Zone aggregate reconciliation to state-level totals",
        "result":f"{pct_pass(zone_ok):.2f}% ({sum(str(r['pass']).lower()=='true' or r['pass'] is True for r in zone_ok)}/{len(zone_ok)})" if zone_ok else "Not available",
        "interpretation":"Uses normalised state names, including NASSARAWA→NASARAWA."
    },
    {
        "dimension":"Precision",
        "measure":"Unexpected non-integer subscription/count cells",
        "result":str(sum(int(r["noninteger_count_cells"]) for r in canon_metrics)) if canon_metrics else "Not available",
        "interpretation":"Subscription and porting counts are expected to be whole-number quantities."
    },
    {
        "dimension":"Traceability",
        "measure":"Canonical sheets with an embedded source/note reference",
        "result":(
            f"{sum(1 for r in metadata_records if latest.get((r['quarter'],r['service_type']),{}).get('workbook')==r['workbook'] and r['embedded_source_or_note_found'])}"
            f"/{sum(1 for r in metadata_records if latest.get((r['quarter'],r['service_type']),{}).get('workbook')==r['workbook'])} canonical sheets"
        ),
        "interpretation":"This is workbook-embedded provenance only; official catalogue provenance must be assessed separately."
    },
    {
        "dimension":"Currentness",
        "measure":"Publication lag",
        "result":"Pending external official catalogue metadata",
        "interpretation":"Local file timestamps are not used."
    },
    {
        "dimension":"Compliance",
        "measure":"Alignment with formal reporting requirements",
        "result":"Pending authoritative NCC/NBS rules",
        "interpretation":"No compliance claim is made without an explicit external benchmark."
    }
]

write_csv(OUT/"v3_01_sheet_occurrences.csv",sheet_occ)
write_csv(OUT/"v3_02_state_sheet_metrics.csv",state_metrics)
write_csv(OUT/"v3_03_canonical_state_data.csv",canonical_state)
write_csv(OUT/"v3_04_cross_period_reconciliation.csv",cross)
write_csv(OUT/"v3_05_porting_reconciliation.csv",porting_records)
write_csv(OUT/"v3_06_zone_reconciliation.csv",zone_recon)
write_csv(OUT/"v3_07_revision_stability.csv",revision)
write_csv(OUT/"v3_08_formula_transparency_drift.csv",formula_drift)
write_csv(OUT/"v3_09_schema_evolution.csv",schema_evo)
write_csv(OUT/"v3_10_metadata_traceability.csv",metadata_records)
write_csv(OUT/"v3_11_issue_register.csv",issues)
write_csv(OUT/"v3_12_dimension_summary.csv",summary)

all_q = sorted({(x["quarter_key"], x["quarter"]) for x in sheet_occ})
actual_quarter_range = (
    f"{all_q[0][1]} to {all_q[-1][1]}" if all_q else "Not available"
)

meta={
    "version":"3.0",
    "workbooks":len(books),
    "discovered_xlsx_files":len(xlsx),
    "skipped_workbooks":skipped_workbooks,
    "quarter_range":actual_quarter_range,
    "key_corrections":[
        "Filename quarter parsing now accepts an underscore immediately before Q (for example Telecoms_Q3_2024.xlsx), preventing silent exclusion of six workbooks.",
        "The audit stops if the expected seven-workbook corpus is not parsed completely.",
        "Traceability denominator is derived from the actual number of canonical sheets rather than hard-coded.",
        "NASSARAWA is normalised to NASARAWA for completeness and zone aggregation, while retained as a label-quality issue.",
        "Voice subtotal/operator rows are detected dynamically rather than from hard-coded row numbers.",
        "Porting operator, total and zone tables are detected dynamically.",
        "Accuracy metrics are reported separately rather than averaged into a single composite percentage.",
        "Workbook-embedded source metadata is distinguished from external catalogue provenance."
    ],
    "tolerances":{"count":COUNT_TOL,"percentage_points":PCT_TOL,"zone":ZONE_TOL}
}
(OUT/"v3_audit_metadata.json").write_text(json.dumps(meta,indent=2),encoding="utf-8")

print("\n"+"="*88)
print("NBS ISO/IEC 25012 AUDIT V3 COMPLETE")
print("="*88)
print("Output:",OUT)
print("\nDimension summary:")
for r in summary:
    print(f"- {r['dimension']} | {r['measure']} | {r['result']}")
print("\nPlease upload:")
print("  v3_12_dimension_summary.csv")
print("  v3_11_issue_register.csv")
print("  v3_04_cross_period_reconciliation.csv")
print("  v3_05_porting_reconciliation.csv")
print("  v3_06_zone_reconciliation.csv")
print("  v3_07_revision_stability.csv")
print("  v3_08_formula_transparency_drift.csv")
print("  v3_09_schema_evolution.csv")